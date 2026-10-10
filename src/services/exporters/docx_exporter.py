"""DOCX 导出器 - 使用 python-docx 生成 Word 文档."""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

from src.utils.logger import get_logger


class DocxExporter:
    """按排版模板与导出选项生成 DOCX 文件."""

    def __init__(self) -> None:
        self.logger = get_logger()

    def export(
        self,
        book_title: str,
        chapters: list,
        template: dict,
        options: dict,
        file_path: Path,
        outline_text: str = "",
        characters_text: str = "",
    ) -> int:
        """执行导出，返回总字数."""
        font = template.get("font") or "SimHei"
        font_size = template.get("font_size") or 14
        line_spacing = template.get("line_spacing") or 1.5
        title_style = template.get("title_style") or "bold_center"
        first_indent = template.get("first_indent", True)
        header_tpl = template.get("header") or ""
        footer_tpl = template.get("footer") or ""

        doc = Document()

        self._set_style_font(doc.styles["Normal"], font, font_size, line_spacing, first_indent)
        for level in (1, 2, 3):
            try:
                self._set_style_font(doc.styles[f"Heading {level}"], font)
            except KeyError:
                pass

        self._setup_header_footer(doc, header_tpl, footer_tpl, book_title)

        total_chars = 0

        if options.get("cover_page", True):
            cover = doc.add_paragraph()
            cover.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = cover.add_run(book_title)
            run.bold = True
            run.font.size = Pt(28)
            doc.add_page_break()

        if options.get("include_outline", False) and outline_text:
            self._add_heading(doc, "大纲", 1, title_style)
            for para in outline_text.split("\n"):
                doc.add_paragraph(para)
            if options.get("chapter_pagination", True):
                doc.add_page_break()

        if options.get("include_characters", False) and characters_text:
            self._add_heading(doc, "角色设定", 1, title_style)
            for para in characters_text.split("\n"):
                doc.add_paragraph(para)
            if options.get("chapter_pagination", True):
                doc.add_page_break()

        current_volume = 0
        first_chapter = True
        for ch in chapters:
            if options.get("chapter_pagination", True) and not first_chapter:
                doc.add_page_break()
            first_chapter = False

            if ch.volume_no != current_volume:
                current_volume = ch.volume_no
                self._add_heading(doc, f"第{current_volume}卷", 1, title_style)

            self._add_heading(doc, f"第{ch.chapter_no}章 {ch.title}", 2, title_style)

            content = ch.content or ""
            for para_text in content.split("\n"):
                doc.add_paragraph(para_text)
            total_chars += len(content)

        doc.save(str(file_path))
        self.logger.info(f"DOCX 导出成功: {file_path}, 共 {total_chars} 字")
        return total_chars

    @staticmethod
    def _set_style_font(
        style,
        font: str,
        font_size: int | None = None,
        line_spacing: float | None = None,
        first_indent: bool | None = None,
    ) -> None:
        style.font.name = font
        if font_size is not None:
            style.font.size = Pt(font_size)
        rpr = style.element.get_or_add_rPr()
        rFonts = rpr.get_or_add_rFonts()
        rFonts.set(qn("w:eastAsia"), font)
        if line_spacing is not None or first_indent is not None:
            pf = style.paragraph_format
            if line_spacing is not None:
                pf.line_spacing = line_spacing
            if first_indent:
                pf.first_line_indent = Pt((font_size or 14) * 2)

    @staticmethod
    def _add_heading(doc, text: str, level: int, title_style: str) -> None:
        heading = doc.add_heading(text, level=level)
        if "center" in (title_style or ""):
            heading.alignment = WD_ALIGN_PARAGRAPH.CENTER

    @staticmethod
    def _setup_header_footer(doc, header_tpl: str, footer_tpl: str, book_title: str) -> None:
        section = doc.sections[0]
        if header_tpl:
            header_para = section.header.paragraphs[0]
            header_para.text = header_tpl.replace("{book_title}", book_title)
        if footer_tpl:
            footer_para = section.footer.paragraphs[0]
            footer_text = footer_tpl.replace("{book_title}", book_title)
            if "{page_number}" in footer_text:
                DocxExporter._add_page_number_field(footer_para, footer_text)
            else:
                footer_para.text = footer_text

    @staticmethod
    def _add_page_number_field(paragraph, template_text: str) -> None:
        parts = template_text.split("{page_number}")
        paragraph.text = parts[0]
        run = paragraph.add_run()
        fld_begin = OxmlElement("w:fldChar")
        fld_begin.set(qn("w:fldCharType"), "begin")
        instr = OxmlElement("w:instrText")
        instr.set(qn("xml:space"), "preserve")
        instr.text = "PAGE"
        fld_end = OxmlElement("w:fldChar")
        fld_end.set(qn("w:fldCharType"), "end")
        run._r.append(fld_begin)
        run._r.append(instr)
        run._r.append(fld_end)
        if len(parts) > 1:
            paragraph.add_run(parts[1])