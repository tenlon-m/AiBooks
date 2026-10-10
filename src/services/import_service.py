"""文件导入服务 - 支持 txt/md/docx/xlsx 四格式解析."""
from __future__ import annotations

from pathlib import Path

from src.api.schemas.import_schema import ImportParseResponse
from src.config.settings import get_config
from src.utils.logger import get_logger

SUPPORTED_FORMATS = {"txt", "md", "docx", "xlsx"}


class ImportService:
    def __init__(self) -> None:
        self.logger = get_logger()
        self.config = get_config()

    def _get_format(self, file_path: str) -> str:
        ext = Path(file_path).suffix.lower().lstrip(".")
        if ext not in SUPPORTED_FORMATS:
            raise ValueError(f"不支持的文件格式: .{ext}，仅支持 {', '.join(SUPPORTED_FORMATS)}")
        return ext

    def _check_size(self, file_path: str) -> None:
        size = Path(file_path).stat().st_size
        max_bytes = self.config.import_max_file_size_mb * 1024 * 1024
        if size > max_bytes:
            raise ValueError(f"文件过大: {size / 1024 / 1024:.1f}MB，上限 {self.config.import_max_file_size_mb}MB")

    def _detect_encoding(self, path: Path) -> str:
        for enc in ("utf-8", "gbk", "gb2312", "big5"):
            try:
                with open(path, "r", encoding=enc) as f:
                    f.read(1024)
                return enc
            except (UnicodeDecodeError, LookupError):
                continue
        return "utf-8"

    def _parse_txt(self, path: Path) -> str:
        encoding = self._detect_encoding(path)
        return path.read_text(encoding=encoding)

    def _parse_md(self, path: Path) -> str:
        import markdown
        html = markdown.markdown(path.read_text(encoding="utf-8"))
        import re
        text = re.sub(r'<[^>]+>', '', html)
        return text.strip()

    def _parse_docx(self, path: Path) -> str:
        from docx import Document
        doc = Document(str(path))
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        return "\n".join(paragraphs)

    def _parse_xlsx(self, path: Path) -> str:
        from openpyxl import load_workbook
        wb = load_workbook(str(path), read_only=True, data_only=True)
        try:
            lines: list[str] = []
            for ws in wb.worksheets:
                for row in ws.iter_rows(values_only=True):
                    cells = [str(c) for c in row if c is not None]
                    if cells:
                        lines.append("\t".join(cells))
            return "\n".join(lines)
        finally:
            wb.close()

    def parse(self, file_path: str, category: str) -> ImportParseResponse:
        path = Path(file_path)
        if not path.exists():
            raise FileNotFoundError(f"文件不存在: {file_path}")
        fmt = self._get_format(file_path)
        self._check_size(file_path)

        try:
            if fmt == "txt":
                content = self._parse_txt(path)
            elif fmt == "md":
                content = self._parse_md(path)
            elif fmt == "docx":
                content = self._parse_docx(path)
            elif fmt == "xlsx":
                content = self._parse_xlsx(path)
            else:
                raise ValueError(f"不支持的格式: {fmt}")
        except Exception as e:
            self.logger.error(f"解析文件失败 {file_path}: {e}")
            raise ValueError(f"解析失败: {e}")

        content = content.strip()
        if not content:
            raise ValueError("文件内容为空")

        return ImportParseResponse(
            content=content,
            char_count=len(content),
            source_format=fmt,
        )